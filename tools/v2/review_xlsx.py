"""Таблица для проверки общего списка слов v2: _scratch/v2/wordlist.csv → review/v2_wordlist.xlsx.

    .venv/Scripts/python tools/v2/review_xlsx.py
"""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "_scratch" / "v2"
OUT = ROOT / "review" / "v2_wordlist.xlsx"
LEVELS = ["a1", "a2", "b1", "b2", "c1", "c2"]
COLORS = {"a1": "D7EFE1", "a2": "D3ECEA", "b1": "DCE5FB", "b2": "E5E1F8", "c1": "F8E0D9", "c2": "F3DEC9"}
KIND = {"main": "слово", "sense": "доп. значение", "phrase": "выражение"}
SOURCE = {
    "oxford": "Oxford 5000", "cefrj": "CEFR-J", "octanove": "Octanove C1/C2", "oxford+cefrj": "Oxford и CEFR-J (среднее)",
    "manual": "вручную", "ielts": "тематика IELTS", "oxford-phrase": "Oxford Phrase List", "awl": "AWL", "nawl": "NAWL",
    "freq": "частотность",
}
POS = {"n": "сущ.", "v": "гл.", "adj": "прил.", "adv": "нареч.", "prep": "предл.", "conj": "союз", "pron": "мест.",
       "det": "опред.", "num": "числ.", "interj": "межд.", "part": "частица", "phr": "выраж.", "phv": "фраз. гл."}
TOPICS = {}  # id → русское название (из ielts_topics.txt)
for line in (Path(__file__).with_name("ielts_topics.txt")).read_text(encoding="utf-8").splitlines():
    if line.startswith("## "):
        tid, _, name = line[3:].partition("|")
        TOPICS[tid.strip()] = name.strip()

HEAD = ["Слово / выражение", "Часть речи", "Тип", "Уровень v2", "Откуда уровень", "Oxford", "CEFR-J", "Octanove",
        "AWL (подсписок)", "Частотность (Zipf)", "Уровень в v1", "Темы IELTS", "Спорное", "Ваша правка"]


def row_of(r: dict) -> list:
    src = r["source"].replace("+freqcap", "")
    note = SOURCE.get(src, src) + (" → снижен по частотности" if "freqcap" in r["source"] else "")
    pos = ", ".join(POS.get(p, p) for p in filter(None, [r["pos"], *r.get("also", "").split(",")]))
    topics = ", ".join(TOPICS.get(t, t) for t in filter(None, r.get("topics", "").split(",")))
    return [r["lemma"], pos, KIND.get(r["card"], r["card"]), r["level"].upper(), note, r["oxford"].upper(), r["cefrj"].upper(),
            r["octanove"].upper(), r["awl"], float(r["zipf"]) if r["zipf"] else None, r["v1"].upper(), topics,
            "да" if r["disputed"] else "", ""]


def sheet(wb: Workbook, title: str, rows: list[list], widths: list[int]) -> None:
    ws = wb.create_sheet(title)
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E4DC")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        lv = (row[3].value or "").lower()
        if lv in COLORS:
            row[3].fill = PatternFill("solid", fgColor=COLORS[lv])
        if row[12].value:
            row[12].fill = PatternFill("solid", fgColor="F6E6A8")
        row[13].fill = PatternFill("solid", fgColor="FFFDF5")


def main() -> None:
    cards = list(csv.DictReader(open(SRC / "wordlist.csv", encoding="utf-8-sig")))
    dropped = list(csv.DictReader(open(SRC / "dropped.csv", encoding="utf-8-sig")))
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    by_level = Counter((c["level"], c["card"]) for c in cards)
    lines = [
        ["English Path 2.0 — общий список слов (на проверку)"],
        [],
        ["Уровень", "Слов", "Доп. значений", "Выражений", "Всего карточек"],
        *[[lv.upper(), by_level[(lv, "main")], by_level[(lv, "sense")], by_level[(lv, "phrase")],
           sum(by_level[(lv, k)] for k in KIND)] for lv in LEVELS],
        ["Итого", sum(v for (l, k), v in by_level.items() if k == "main"), sum(v for (l, k), v in by_level.items() if k == "sense"),
         sum(v for (l, k), v in by_level.items() if k == "phrase"), len(cards)],
        [],
        ["Как определялся уровень"],
        ["1. Oxford 3000/5000 (A1–C1) — главный ориентир."],
        ["2. Если Oxford и CEFR-J расходятся на 2+ уровня — средний уровень с округлением вниз (лист «Спорные»)."],
        ["3. Нет в Oxford — CEFR-J (A1–B2), затем Octanove (C1–C2)."],
        ["4. Очень частотное слово не может быть «продвинутым»: потолок по частотности (Zipf ≥ 5.5 → не выше A2, ≥ 5.0 → B1, ≥ 4.6 → B2)."],
        ["5. Слова вне списков (из v1, AWL/NAWL) — уровень назначен вручную по смыслу."],
        ["6. Тематика IELTS (22 темы), идиомы и фразовые глаголы — добавлены отдельным слоем B1–C2."],
        ["7. Исключены редкие производные (-ly, -ness, un-…), обрывки и узкие термины — лист «Исключено»."],
        [],
        ["Многозначные слова: основное значение — на своём уровне; другая часть речи на более высоком уровне — отдельная"],
        ["карточка «доп. значение» (book сущ. A1 / book гл. A2). Значения внутри одной части речи будут разобраны при наполнении."],
        [],
        ["Что проверить: уровень слова (колонка «Уровень v2»), лишние и недостающие слова. Правки — в колонке «Ваша правка»:"],
        ["например «A2», «убрать», «добавить: tree house». Можно просто написать комментарий."],
        [],
        ["Источники: Oxford 3000/5000 и Oxford Phrase List (Oxford University Press); CEFR-J Vocabulary Profile 1.5 (Tono Laboratory, TUFS);"],
        ["Octanove Vocabulary Profile C1/C2 (Octanove Labs, CC BY-SA 4.0); AWL (Coxhead, 2000); NAWL (Browne, Culligan, Phillips);"],
        ["wordfreq (частотность, Speer et al.); IELTS Vocabulary Booster (N. Miles) — сверка тематической лексики."],
    ]
    for ln in lines:
        ws.append(ln)
    ws["A1"].font = Font(bold=True, size=14)
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in (10, 18):
        ws.cell(r, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    for col in "BCDE":
        ws.column_dimensions[col].width = 16

    widths = [28, 16, 15, 11, 30, 9, 9, 10, 10, 12, 11, 30, 10, 26]
    order = {"main": 0, "sense": 1, "phrase": 2}
    for lv in LEVELS:
        rs = sorted((c for c in cards if c["level"] == lv), key=lambda c: (order[c["card"]], c["lemma"].lower()))
        sheet(wb, lv.upper(), [row_of(r) for r in rs], widths)
    disputed = sorted((c for c in cards if c["disputed"]), key=lambda c: (LEVELS.index(c["level"]), c["lemma"].lower()))
    sheet(wb, "Спорные", [row_of(r) for r in disputed], widths)
    ws = wb.create_sheet("Исключено")
    ws.append(["Слово", "Часть речи", "Причина", "Oxford", "CEFR-J", "Octanove", "Частотность (Zipf)", "Уровень в v1"])
    for d in dropped:
        ws.append([d["lemma"], POS.get(d["pos"], d["pos"]), d["dropped"], d["oxford"].upper(), d["cefrj"].upper(),
                   d["octanove"].upper(), float(d["zipf"]), d["v1"].upper()])
    for i, w in enumerate([26, 12, 28, 9, 9, 10, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(bold=True)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("сохранено:", OUT)


if __name__ == "__main__":
    main()
