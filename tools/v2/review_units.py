"""Таблица уроков v2 на проверку: _scratch/v2/units.csv → review/v2_units.xlsx."""

import csv
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).parent))
from topics import TOPICS  # noqa: E402

LEVELS = ["a1", "a2", "b1", "b2", "c1", "c2"]
units = list(csv.DictReader(open(ROOT / "_scratch" / "v2" / "units.csv", encoding="utf-8-sig")))
wb = Workbook()
ws = wb.active
ws.title = "Сводка"
ws.append(["English Path 2.0 — темы и уроки (на проверку)"])
ws["A1"].font = Font(bold=True, size=14)
ws.append([])
ws.append(["Тема", *[lv.upper() for lv in LEVELS], "Всего уроков"])
for c in ws[3]:
    c.font = Font(bold=True)
count = Counter((u["topic"], u["level"]) for u in units)
for tid, (ru, _) in TOPICS.items():
    row = [count[(tid, lv)] for lv in LEVELS]
    if sum(row):
        ws.append([ru, *[r or "" for r in row], sum(row)])
ws.append(["Итого", *[sum(1 for u in units if u["level"] == lv) for lv in LEVELS], len(units)])
ws.append([])
ws.append(["Урок — 15–25 карточек одной темы: похожие по смыслу слова рядом, частотные раньше."])
ws.append(["Порядок (№): тематические уроки в учебной последовательности, общие (фразы, служебные слова, описание) — вперемешку между ними."])
ws.append(["Если тема на уровне слишком мала (< 8 карточек), она объединена с соседней."])
ws.append(["Правки — в колонке «Ваша правка» на листах уровней: перенести слово, переименовать тему, разделить урок."])
ws.column_dimensions["A"].width = 36
for lv in LEVELS:
    s = wb.create_sheet(lv.upper())
    s.append(["№", "Тема", "Урок темы", "Слов", "Слова урока", "Ваша правка"])
    for c in s[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E4DC")
    for u in sorted((u for u in units if u["level"] == lv), key=lambda u: int(u["seq"])):
        s.append([int(u["seq"]), u["titleRu"], int(u["n"]), int(u["size"]), u["words"], ""])
    for col, w in zip("ABCDEF", (5, 40, 8, 6, 140, 30)):
        s.column_dimensions[col].width = w
    for row in s.iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    s.freeze_panes = "C2"
    s.auto_filter.ref = s.dimensions
out = ROOT / "review" / "v2_units.xlsx"
wb.save(out)
print("сохранено", out.name, len(units), "уроков")
